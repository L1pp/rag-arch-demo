"""Opt-in local Linux/LibreOffice checks; no database, model or OCR calls."""
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest

from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.testclient import TestClient
import pymupdf

from rag_arch.document_pages import DocumentPageService
from rag_arch.loaders import load_document
from rag_arch.office_conversion import OfficeConverter
from rag_arch.web.document_routes import create_document_router
from rag_arch.web.documents import DocumentStore


ENABLED = sys.platform.startswith("linux") and os.environ.get("RAG_ARCH_TEST_LIBREOFFICE") == "1"


@unittest.skipUnless(ENABLED, "requires local Linux LibreOffice and explicit integration opt-in")
class LinuxOfficeIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary_parent = Path(os.environ.get("RAG_ARCH_LINUX_TEST_ROOT", tempfile.gettempdir()))
        self.temporary_parent.mkdir(parents=True, mode=0o700, exist_ok=True)
        self.root = Path(tempfile.mkdtemp(prefix="rag-arch-office-integration-", dir=self.temporary_parent))
        self.addCleanup(self._cleanup)
        self.output = os.environ.get("RAG_ARCH_LINUX_TEST_ARTIFACT_DIR")

    def _cleanup(self):
        marker = self.root / "data/runtime/office_conversion/active.json"
        if marker.exists():
            if self.output:
                diagnostic = Path(self.output) / (self.root.name + "-diagnostics")
                diagnostic.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(marker, diagnostic / "active.json")
                owner = json.loads(marker.read_text(encoding="utf-8"))
                run = Path(owner.get("run_directory", ""))
                if run.is_relative_to(self.root):
                    for name in ("result.json", "children.json", "converter.log", "cancel.json"):
                        source = run / name
                        if source.is_file():
                            shutil.copyfile(source, diagnostic / name)
            # An uncertain conversion owns its files until recovery confirms exit.
            print("Preserving uncertain integration directory:", self.root)
            return
        if self.root.name.startswith("rag-arch-office-integration-") and self.root.parent == self.temporary_parent:
            shutil.rmtree(self.root)

    def make_word(self):
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Noto Sans CJK SC"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        document.add_paragraph("LINUX_PAGE_ONE Linux兼容验证")
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = "结构参数"
        cells[1].text = "42 mm / 1‰"
        document.add_page_break()
        document.add_paragraph("LINUX_PAGE_TWO 原页图补充")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def upload_and_verify(self, data, filename):
        store = DocumentStore(self.root, "integration")
        self.addCleanup(store.close)
        app = FastAPI()
        app.state.documents = store
        app.state.tasks = SimpleNamespace(workspace_id="integration", user_id=None)
        app.include_router(create_document_router())
        with TestClient(app) as client:
            uploaded = client.post("/api/uploads", params={"workspace_id": "integration", "filename": filename}, content=data)
            self.assertEqual(uploaded.status_code, 201, uploaded.text)
            asset_id = uploaded.json()["asset_id"]
            asset = store.get(asset_id)
            original_stamp = asset.path.stat().st_mtime_ns
            original = client.get(f"/api/uploads/{asset_id}/download", params={"workspace_id": "integration"})
            self.assertEqual(original.content, data)
            loaded = store.load(asset_id)
            self.assertEqual(loaded.filename, filename)
            self.assertEqual(loaded.source_kind, Path(filename).suffix[1:])
            text = "\n".join(unit.text for unit in loaded.units)
            self.assertIn("LINUX_PAGE_ONE", text)
            self.assertIn("LINUX_PAGE_TWO", text)
            self.assertIn("结构参数", text)
            self.assertIn("42", text)
            info = client.get("/api/documents/pages", params={"workspace_id": "integration", "asset": asset_id, "page": 2})
            self.assertEqual(info.status_code, 200, info.text)
            self.assertEqual(info.json()["page_count"], 2)
            image = client.get(info.json()["pages"][0]["image_url"])
            self.assertEqual(image.status_code, 200, image.text)
            self.assertTrue(image.content.startswith(b"\x89PNG\r\n\x1a\n"))
            shared_pdf = asset.path.with_name("pages.pdf")
            self.assertTrue(shared_pdf.is_file())
            with pymupdf.open(shared_pdf) as pdf:
                self.assertIn("LINUX_PAGE_TWO", pdf[1].get_text())
            before = shared_pdf.stat().st_mtime_ns
            # A new service must reuse exactly the PDF already used by the API.
            service = DocumentPageService(self.root)
            normalized = asset.path.with_name("normalized.docx") if filename.endswith(".doc") else None
            service.register("same-source", filename, path=asset.path, normalized_docx_path=normalized)
            self.assertIn("LINUX_PAGE_TWO", service.read_page("same-source", 2)["text"])
            self.assertEqual(shared_pdf.stat().st_mtime_ns, before)
            self.assertEqual(asset.path.read_bytes(), data)
            self.assertEqual(asset.path.stat().st_mtime_ns, original_stamp)
            self.assertFalse((self.root / "data/runtime/office_conversion/active.json").exists())
            if self.output:
                output = Path(self.output)
                output.mkdir(parents=True, exist_ok=True)
                (output / (Path(filename).suffix[1:] + "-page2.png")).write_bytes(image.content)
                shutil.copyfile(shared_pdf, output / (Path(filename).suffix[1:] + "-pages.pdf"))
                (output / (Path(filename).suffix[1:] + "-result.json")).write_text(json.dumps({
                    "filename": filename, "source_kind": loaded.source_kind, "pages": 2,
                    "original_unchanged": True, "shared_pdf_reused": True,
                    "online_ocr_calls": 0, "model_calls": 0,
                }, ensure_ascii=False, indent=2), encoding="utf-8")

    def test_docx_upload_text_and_cached_pages(self):
        self.upload_and_verify(self.make_word(), "兼容方案.docx")

    def test_real_legacy_doc_upload_text_and_cached_pages(self):
        # Generate a genuine old-format DOC from our own harmless synthetic file.
        fixtures = self.root / "fixtures"
        fixtures.mkdir()
        source = fixtures / "fixture.docx"
        source.write_bytes(self.make_word())
        profile = self.root / "fixture-profile"
        executable = OfficeConverter(self.root)._executable()
        generated = subprocess.run([
            executable, "-env:UserInstallation=" + profile.as_uri(), "--headless",
            "--nologo", "--nodefault", "--norestore", "--convert-to", "doc:MS Word 97",
            "--outdir", str(fixtures), str(source),
        ], capture_output=True, timeout=60)
        legacy = fixtures / "fixture.doc"
        self.assertEqual(generated.returncode, 0, generated.stderr.decode(errors="replace"))
        self.assertTrue(legacy.is_file(), generated.stdout.decode(errors="replace"))
        data = legacy.read_bytes()
        self.assertTrue(data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"))
        self.upload_and_verify(data, "兼容方案.doc")

    def test_image_only_page_stays_available_without_online_ocr(self):
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=300, height=400)
            pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), False)
            pixmap.clear_with(180)
            page.insert_image(page.rect, stream=pixmap.tobytes("png"))
            data = pdf.tobytes()
        loaded = load_document("scan.pdf", data)
        self.assertEqual(loaded.units, [])
        pages = DocumentPageService(self.root)
        pages.register("scan", "scan.pdf", data=data)
        page = pages.read_page("scan", 1)
        self.assertEqual(page["text"], "")
        self.assertTrue(page["image_bytes"].startswith(b"\x89PNG"))


if __name__ == "__main__":
    unittest.main()
